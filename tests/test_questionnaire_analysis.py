from pathlib import Path

import pytest
from openpyxl import load_workbook

from crq.analysis import gap_analysis, load_case, packages
from crq.cli import main
from crq.controls import patch_exposure_reduction
from crq.excel import write_workbook
from crq.questionnaire import assess, best_answers, load_questions, posture_from_answers

ROOT = Path(__file__).resolve().parents[1]
PROFILES = sorted((ROOT / "profiles").glob("*.yaml"))


def test_best_answers_score_100_and_preferred():
    a = assess(best_answers())
    assert a.score == 100.0 and a.tier == "Preferred" and a.critical_failures == []
    assert all(v == 1.0 for v in a.posture.as_dict().values() if v != a.posture.level("patching"))
    assert a.posture.level("patching") == pytest.approx(patch_exposure_reduction(7))


def test_knockouts_force_referral_regardless_of_score():
    ans = best_answers()
    ans["backups"] = "offsite_online"
    a = assess(ans)
    assert a.score > 85 and a.tier == "Refer / restricted" and a.critical_failures == ["backups"]


def test_eol_modifier_reduces_patching_level():
    ans = best_answers()
    base = posture_from_answers(ans).level("patching")
    ans["eol_systems"] = "unsupported"
    assert posture_from_answers(ans).level("patching") == pytest.approx(base * 0.7)


def test_bad_answers_rejected():
    ans = best_answers()
    ans["mfa_email"] = "maybe"
    with pytest.raises(ValueError):
        assess(ans)
    del ans["mfa_email"]
    with pytest.raises(ValueError):
        assess(ans)


@pytest.mark.parametrize("path", PROFILES, ids=lambda p: p.stem)
def test_profiles_load_and_answer_every_question(path):
    case = load_case(path)
    assert set(case.answers) == set(load_questions())
    assert "hypothetical" in case.profile.name.lower()


def test_gap_values_nonnegative_and_packages_ordered():
    case = load_case(ROOT / "profiles" / "dental_clinic.yaml")
    a, base, rows = gap_analysis(case, n=20_000, seed=1)
    modelled = [r for r in rows if r.modelled]
    assert modelled and all(r.delta_ale >= -1e-6 for r in modelled)
    assert rows[0].critical  # knock-outs listed first
    pk = packages(case, n=20_000, seed=1)
    assert pk["current"][1].ale >= pk["fix_knockouts"][1].ale >= pk["all_controls"][1].ale
    assert pk["all_controls"][0].tier == "Preferred"


def test_excel_workbook(tmp_path):
    case = load_case(ROOT / "profiles" / "law_firm.yaml")
    a, base, rows = gap_analysis(case, n=10_000, seed=2)
    pk = packages(case, n=10_000, seed=2)
    path = write_workbook(case, a, base, rows, pk, tmp_path / "wb.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Questionnaire", "Gaps & premium impact", "Evidence checklist", "Loss exceedance",
                             "Assumptions & sources", "Bibliography"]
    assert "HYPOTHETICAL" in wb["Summary"]["A2"].value
    assert wb["Assumptions & sources"].max_row > 60


def test_cli_run_and_assess(capsys, tmp_path):
    assert main(["assess", str(ROOT / "profiles" / "manufacturer.yaml")]) == 0
    assert "KNOCK-OUT" in capsys.readouterr().out
    out = tmp_path / "m.xlsx"
    assert main(["run", str(ROOT / "profiles" / "manufacturer.yaml"), "--n", "5000", "--xlsx", str(out)]) == 0
    assert out.exists() and "ALE" in capsys.readouterr().out
