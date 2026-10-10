"""Underwriter workbook: one .xlsx per firm with summary, gaps, evidence and assumptions."""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analysis import Case, GapRow, indicated_premium
from .model import SCENARIO_LABELS, SimulationResult
from .parameters import ParameterSet, load_sources
from .questionnaire import Assessment, load_questions

HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
MONEY = '"$"#,##0'
PCT = "0.00%"
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row: int, cols: list[str]) -> None:
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP


def _widths(ws, widths: list[int]) -> None:
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def write_workbook(case: Case, assessment: Assessment, res: SimulationResult, gaps: list[GapRow],
                   packages: dict[str, tuple[Assessment, SimulationResult]], path: str | Path,
                   params: ParameterSet | None = None) -> Path:
    params = params or ParameterSet()
    prof = case.profile
    wb = Workbook()

    # --- Summary -------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Cyber underwriting readiness: {prof.name}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("HYPOTHETICAL firm. Model output is illustrative, built from public data and stated assumptions; "
                "it is not a rate, a quote, or a coverage opinion.")
    ws["A2"].font = Font(italic=True, color="9C0006")
    rows = [
        ("Industry", prof.industry), ("Employees", prof.employees), ("Annual revenue", prof.annual_revenue),
        ("Records held", prof.records_held), ("Policy aggregate limit", prof.policy.aggregate_limit),
        ("Retention", prof.policy.retention), ("Social-engineering sublimit", prof.policy.social_engineering_sublimit),
        ("Readiness score (0-100)", assessment.score), ("Tier", assessment.tier), ("Tier meaning", assessment.tier_meaning),
        ("Knock-out gaps", ", ".join(assessment.critical_failures) or "none"),
        ("Simulated years", res.n), ("Random seed", res.seed),
    ]
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, float) and v > 1000:
            c.number_format = MONEY
    r0 = len(rows) + 6
    _header(ws, r0, ["Metric", "Current", "Fix knock-outs", "All controls"])
    lr = params.central("pricing.loss_ratio")
    metrics = [
        ("Annualized loss expectancy (ALE)", lambda r: r.ale, MONEY),
        ("P(any loss event in a year)", lambda r: r.prob_any_event(), PCT),
        ("P(annual loss > $100K)", lambda r: float(r.exceedance([1e5])[0]), PCT),
        ("P(annual loss > $1M)", lambda r: float(r.exceedance([1e6])[0]), PCT),
        ("1-in-100-year loss (P99)", lambda r: r.quantile(0.99), MONEY),
        ("TVaR 99% (mean of worst 1% of years)", lambda r: r.tvar(0.99), MONEY),
        ("Expected insured loss", lambda r: float(r.insured_loss.mean()), MONEY),
        (f"Indicated premium (expected insured loss / {lr:.0%} loss ratio)", lambda r: indicated_premium(r, params), MONEY),
    ]
    for i, (name, fn, fmt) in enumerate(metrics, r0 + 1):
        ws.cell(row=i, column=1, value=name)
        for j, key in enumerate(("current", "fix_knockouts", "all_controls"), 2):
            c = ws.cell(row=i, column=j, value=float(fn(packages[key][1])))
            c.number_format = fmt
    r1 = r0 + len(metrics) + 2
    _header(ws, r1, ["Scenario", "ALE (current)", "Mean events / yr", "Share of ALE"])
    for i, (s, v) in enumerate(res.ale_by_scenario().items(), r1 + 1):
        ws.cell(row=i, column=1, value=SCENARIO_LABELS[s])
        ws.cell(row=i, column=2, value=v).number_format = MONEY
        ws.cell(row=i, column=3, value=float(res.lam[s].mean())).number_format = "0.0000"
        ws.cell(row=i, column=4, value=v / res.ale if res.ale else 0).number_format = PCT
    _widths(ws, [52, 18, 18, 18])

    # --- Questionnaire -------------------------------------------------------
    qs = load_questions()
    wq = wb.create_sheet("Questionnaire")
    _header(wq, 1, ["Section", "Question", "Answer", "Score (0-1)", "Weight", "Knock-out", "Modelled"])
    for i, (qid, q) in enumerate(qs.items(), 2):
        vals = [q.section, q.text, q.options[assessment.answers[qid]]["label"], assessment.question_scores[qid], q.weight,
                "yes" if q.critical else "", "yes" if q.modelled else "scored only"]
        for j, v in enumerate(vals, 1):
            wq.cell(row=i, column=j, value=v).alignment = WRAP
    _widths(wq, [22, 70, 45, 12, 9, 10, 12])

    # --- Gaps ----------------------------------------------------------------
    wg = wb.create_sheet("Gaps & premium impact")
    cols = ["Control", "Knock-out", "Current answer", "Recommended", "ALE reduction", "Reduction in TVaR99 (mean of worst 1% of years)", "Reduction in P(annual loss > $100K)",
            "Change in indicated premium", "Illustrative annual cost", "ALE reduction per $ spent", "Evidence strength",
            "Underwriting / premium narrative"]
    _header(wg, 1, cols)
    for i, g in enumerate(gaps, 2):
        vals = [g.question, "YES" if g.critical else "", g.current, g.recommended,
                None if np.isnan(g.delta_ale) else g.delta_ale, None if np.isnan(g.delta_tvar99) else g.delta_tvar99,
                None if np.isnan(g.delta_p_over_100k) else g.delta_p_over_100k,
                None if np.isnan(g.premium_change_pct) else g.premium_change_pct / 100, g.annual_cost,
                None if (np.isnan(g.delta_ale) or not g.annual_cost) else g.delta_ale / g.annual_cost,
                g.evidence_strength, g.note if g.modelled else "Scored but not simulated (no defensible public effect size). " + g.note]
        for j, v in enumerate(vals, 1):
            c = wg.cell(row=i, column=j, value=v)
            c.alignment = WRAP
            if j in (5, 6, 9):
                c.number_format = MONEY
            if j == 7:
                c.number_format = "0.000%"
            if j == 8:
                c.number_format = "0.0%"
            if j == 10:
                c.number_format = "0.00"
    wg.cell(row=len(gaps) + 3, column=1,
            value="Gap values are one-at-a-time (common random numbers); they do not add up because controls overlap. "
                  
                  "Premium change = change in expected insured loss (loss-cost view); real pricing also reflects filed rates, minimums and market cycle.")
    _widths(wg, [20, 10, 34, 34, 14, 16, 14, 14, 14, 12, 12, 80])

    # --- Evidence checklist --------------------------------------------------
    we = wb.create_sheet("Evidence checklist")
    _header(we, 1, ["Control", "Status", "Evidence to collect for the application / renewal"])
    row = 2
    for qid, q in qs.items():
        status = "meets best practice" if assessment.question_scores[qid] >= q.options[q.best_option]["score"] else "GAP"
        for ev in q.evidence:
            for j, v in enumerate([qid, status, ev], 1):
                we.cell(row=row, column=j, value=v).alignment = WRAP
            row += 1
    _widths(we, [22, 20, 100])

    # --- Loss distribution ---------------------------------------------------
    wd = wb.create_sheet("Loss exceedance")
    _header(wd, 1, ["Annual loss threshold", "P(exceed) current", "P(exceed) fix knock-outs", "P(exceed) all controls"])
    thresholds = [1e3, 5e3, 1e4, 2.5e4, 5e4, 1e5, 2.5e5, 5e5, 1e6, 2.5e6, 5e6]
    for i, t in enumerate(thresholds, 2):
        wd.cell(row=i, column=1, value=t).number_format = MONEY
        for j, key in enumerate(("current", "fix_knockouts", "all_controls"), 2):
            wd.cell(row=i, column=j, value=float(packages[key][1].exceedance([t])[0])).number_format = "0.000%"
    _widths(wd, [22, 20, 24, 22])

    # --- Assumptions ---------------------------------------------------------
    wa = wb.create_sheet("Assumptions & sources")
    _header(wa, 1, ["Parameter", "Description", "Distribution", "Basis", "Sources", "Note"])
    for i, p in enumerate(params.registry.values(), 2):
        for j, v in enumerate([p.name, p.description, p.describe_range(), p.basis, ", ".join(p.sources), p.note], 1):
            wa.cell(row=i, column=j, value=v).alignment = WRAP
    _widths(wa, [30, 50, 28, 12, 26, 70])
    srcs = load_sources()
    ws2 = wb.create_sheet("Bibliography")
    _header(ws2, 1, ["Id", "Title", "Publisher", "Year", "URL", "Used for"])
    for i, (sid, s) in enumerate(srcs.items(), 2):
        for j, v in enumerate([sid, s["title"], s["publisher"], s["year"], s["url"], s["used_for"]], 1):
            ws2.cell(row=i, column=j, value=v).alignment = WRAP
    _widths(ws2, [18, 45, 30, 8, 50, 80])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
